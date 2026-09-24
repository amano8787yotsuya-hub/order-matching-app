import io
import os
import re
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
import pandas as pd
import pypdf
import streamlit as st

st.set_page_config(page_title="発注確定版 自動生成ツール", layout="wide")
st.title("📦 発注予定表 × ピッキング確定 突合生成システム")
st.caption(
    "予定表とPDFを読み込み、未登録商品・店舗の自動追加、確定数量反映、マクロ内蔵の確定版（.xlsm）を出力します。"
)

YELLOW_FILL = PatternFill(
    start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"
)
MEIRYO_FONT = Font(name="Meiryo UI", size=10)
MEIRYO_HEADER_FONT = Font(name="Meiryo UI", size=10, bold=True)
DIFF_NUM_FORMAT = '#,##0;[Red]-#,##0;""'


# -------------------------------------------------------------
# 1. PDF解析関数 (アピタ・ドンキ共通)
# -------------------------------------------------------------
def parse_picking_pdf(file_bytes):
  reader = pypdf.PdfReader(io.BytesIO(file_bytes))
  records = []
  current_store_code = None
  current_store_name = None

  for page in reader.pages:
    text = page.extract_text()
    if not text:
      continue
    lines = [line.strip() for line in text.split("\n") if line.strip()]
    i = 0
    while i < len(lines):
      line = lines[i]
      if "納入先：" in line:
        m = re.search(r"納入先：\s*(\d+)\s*(.*)", line)
        if m:
          current_store_code = str(int(m.group(1)))
          current_store_name = m.group(2).strip()

      m_item = re.match(r"^(\d{8})\s+(.*)", line)
      if m_item and current_store_code:
        item_code = str(int(m_item.group(1)))
        item_name = m_item.group(2).strip()
        price, qty = None, None
        for step in range(1, 4):
          if i + step < len(lines):
            nxt = lines[i + step]
            m_price = re.search(r"[ー-]\s*\d+\s*([\d,]+)", nxt)
            if m_price and price is None:
              price = int(m_price.group(1).replace(",", ""))
            elif nxt.isdigit() and qty is None:
              qty = int(nxt)
        if qty is not None:
          records.append({
              "store_code": current_store_code,
              "store_name": current_store_name,
              "item_code": item_code,
              "item_name": item_name,
              "price": price if price else 0,
              "qty": qty,
          })
      i += 1
  return pd.DataFrame(records)


# -------------------------------------------------------------
# 2. 列幅の自動調整
# -------------------------------------------------------------
def auto_fit_columns(ws, max_cols=6):
  for col in range(1, max_cols + 1):
    col_letter = get_column_letter(col)
    max_len = 0
    for r in range(1, min(ws.max_row + 1, 100)):
      val = ws.cell(r, col).value
      if val is not None:
        s_val = str(val)
        val_len = sum(2 if ord(ch) > 127 else 1 for ch in s_val)
        if val_len > max_len:
          max_len = val_len
    ws.column_dimensions[col_letter].width = max(max_len + 3, 11)


# -------------------------------------------------------------
# 3. メイン処理（突合 ＆ エクセル再構築）
# -------------------------------------------------------------
def process_data(excel_file, apita_pdf, donki_pdf):
  xls_all = pd.read_excel(excel_file, sheet_name=None, header=None)

  # template.xlsm（マクロ有効ブック）をベースに読み込み
  if os.path.exists("template.xlsm"):
    wb = openpyxl.load_workbook("template.xlsm", keep_vba=True)
  else:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

  # 【改善1】既存シートのゴミを残さないよう、完全にシートを再作成
  for sname in list(wb.sheetnames):
    if sname not in ["ルート", "店舗マスタ"]:
      del wb[sname]

  date_sheet_name = None
  for s_name, s_df in xls_all.items():
    if s_name not in ["ルート", "店舗マスタ"] and date_sheet_name is None:
      date_sheet_name = s_name

    # 新規作成して1行目から書き込む（appendによる行番号のズレを根絶）
    ws_tmp = wb.create_sheet(title=s_name)
    for r_idx, r_vals in enumerate(s_df.itertuples(index=False), start=1):
      for c_idx, val in enumerate(r_vals, start=1):
        if pd.notna(val):
          ws_tmp.cell(r_idx, c_idx, val)

  ws = wb[date_sheet_name]

  # 納品日の判定
  raw_date_val = ws.cell(1, 2).value
  if not raw_date_val:
    raw_date_val = ws.cell(1, 3).value

  date_clean = str(raw_date_val)
  if " " in date_clean:
    date_clean = date_clean.split(" ")[0]
  elif " " in date_clean:
    date_clean = date_clean.split(" ")[0]

  try:
    target_dt = pd.to_datetime(date_clean)
    is_thursday = target_dt.weekday() == 3
  except:
    is_thursday = False

  # ルートシートから配送区分を取得
  dist_map = {}
  ws_route = None
  for r_name in ["ルート", "店舗マスタ"]:
    if r_name in wb.sheetnames:
      ws_route = wb[r_name]
      break

  if ws_route:
    c_code = 4 if is_thursday else 1
    c_dist = 6 if is_thursday else 3
    for r in range(1, ws_route.max_row + 1):
      sc_val = ws_route.cell(r, c_code).value
      dist_val = ws_route.cell(r, c_dist).value
      if sc_val is not None and str(sc_val).strip() != "":
        try:
          sc_k = str(int(float(str(sc_val).strip())))
          dist_str = str(dist_val).strip() if dist_val else "自社"
          dist_map[sc_k] = "佐川便のみ" if "佐川" in dist_str else "自社便のみ"
        except:
          pass

  # A列に「採用」フラグがなければ挿入
  if ws.cell(2, 1).value != "採用":
    ws.insert_cols(1)
    ws.cell(2, 1).value = "採用"
    for r in range(3, ws.max_row + 1):
      if (
          ws.cell(r, 2).value is not None
          and str(ws.cell(r, 2).value).strip() != ""
      ):
        ws.cell(r, 1).value = True

  ws.cell(1, 2).value = raw_date_val
  ws.freeze_panes = "E3"

  # PDFの解析とマージ
  df_apita = parse_picking_pdf(apita_pdf.read())
  df_donki = parse_picking_pdf(donki_pdf.read())
  df_pdf = pd.concat([df_apita, df_donki], ignore_index=True)
  df_pdf = df_pdf.drop_duplicates(
      subset=["store_code", "item_code"], keep="last"
  )

  pdf_confirmed_stores = set(df_pdf["store_code"].unique())
  pdf_qty_dict = {
      (row["store_code"], row["item_code"]): row["qty"]
      for _, row in df_pdf.iterrows()
  }

  # 店舗列のマッピング (E列: col 5 〜)
  store_col_map = {}
  total_col_idx = None
  for c in range(5, ws.max_column + 1):
    v1 = ws.cell(1, c).value
    v2 = ws.cell(2, c).value
    if str(v2).strip() == "総計":
      total_col_idx = c
      break
    if v1 is not None and str(v1).strip() != "":
      try:
        sc = str(int(float(str(v1).strip())))
        store_col_map[sc] = c
      except:
        pass

  if total_col_idx is None:
    total_col_idx = ws.max_column + 1
    ws.cell(2, total_col_idx).value = "総計"

  # 商品行のマッピング
  item_row_map = {}
  orig_totals = {}
  for r in range(3, ws.max_row + 1):
    c_val = ws.cell(r, 2).value
    if c_val is not None and str(c_val).strip() != "":
      try:
        ic = str(int(float(str(c_val).strip())))
        item_row_map[ic] = r
        row_sum = sum(
            int(ws.cell(r, c).value)
            for c in range(5, total_col_idx)
            if ws.cell(r, c).value is not None
            and str(ws.cell(r, c).value).isdigit()
        )
        orig_totals[ic] = row_sum
      except:
        pass

  # 未登録店舗の追加（総計列の前に追加）
  added_stores_list = []
  pdf_stores = df_pdf[["store_code", "store_name"]].drop_duplicates()
  for _, s_row in pdf_stores.iterrows():
    sc = s_row["store_code"]
    if sc not in store_col_map:
      ws.insert_cols(total_col_idx)
      new_col = total_col_idx
      ws.cell(1, new_col).value = sc
      ws.cell(2, new_col).value = s_row["store_name"]
      store_col_map[sc] = new_col
      added_stores_list.append(f"{sc} ({s_row['store_name']})")
      total_col_idx += 1

  # 未登録商品の追加（最終行に追加）
  added_items_list = []
  pdf_items = df_pdf[["item_code", "item_name", "price"]].drop_duplicates(
      subset=["item_code"]
  )
  current_last_row = ws.max_row
  for _, i_row in pdf_items.iterrows():
    ic = i_row["item_code"]
    if ic not in item_row_map:
      current_last_row += 1
      ws.cell(current_last_row, 1).value = True
      ws.cell(current_last_row, 2).value = ic
      ws.cell(current_last_row, 3).value = i_row["item_name"]
      ws.cell(current_last_row, 4).value = i_row["price"]

      # 【改善2】新商品行の店舗数量セル（E列〜総計前列）を確実にクリア
      for cl in range(5, total_col_idx):
        ws.cell(current_last_row, cl).value = None

      item_row_map[ic] = current_last_row
      orig_totals[ic] = 0
      added_items_list.append(f"{ic} {i_row['item_name']}")

  last_store_letter = get_column_letter(total_col_idx - 1)
  for r in range(3, current_last_row + 1):
    ws.cell(r, total_col_idx).value = (
        f"=IF(AND(A{r}=TRUE,SUM(E{r}:{last_store_letter}{r})>0),SUM(E{r}:{last_store_letter}{r}),"")"
    )

  ws.auto_filter.ref = f"A2:{last_store_letter}{current_last_row}"

  # 差分の突合 ＆ 反映
  log_info = []
  diff_count = 0
  for sc, c_idx in store_col_map.items():
    if sc in pdf_confirmed_stores:
      for ic, r_idx in item_row_map.items():
        old_val = ws.cell(r_idx, c_idx).value
        old_qty = (
            int(old_val)
            if (old_val is not None and str(old_val).isdigit())
            else 0
        )
        new_qty = pdf_qty_dict.get((sc, ic), 0)

        if old_qty != new_qty:
          target_cell = ws.cell(r_idx, c_idx)
          target_cell.value = new_qty if new_qty > 0 else None
          target_cell.fill = YELLOW_FILL
          diff_count += 1

          log_info.append({
              "sheet": date_sheet_name,
              "store_code": sc,
              "store_name": ws.cell(2, c_idx).value,
              "item_code": ws.cell(r_idx, 2).value,
              "item_name": ws.cell(r_idx, 3).value,
              "item_row": r_idx,
              "store_col": c_idx,
              "store_col_letter": get_column_letter(c_idx),
              "old_qty": old_qty,
              "new_qty": new_qty,
              "diff": new_qty - old_qty,
          })

  # -------------------------------------------------------------
  # 【改善3】修正差分ログシートの再構築（完全数式リンク ＆ 並び替え耐性）
  # -------------------------------------------------------------
  if "修正差分ログ" in wb.sheetnames:
    del wb["修正差分ログ"]
  ws_log = wb.create_sheet(title="修正差分ログ")
  headers_log = [
      "シート",
      "店舗コード",
      "店舗名",
      "商品コード",
      "商品名",
      "修正前(予定)",
      "修正後(確定)",
      "増減",
  ]
  ws_log.append(headers_log)

  for log_idx, item in enumerate(log_info, start=2):
    r_no = item["item_row"]
    c_let = item["store_col_letter"]
    sc_val = item["store_code"]
    st_name = item["store_name"]

    # 日付シートの採用フラグ連動
    ws_log.cell(
        log_idx, 1
    ).value = f'=IF(\'{date_sheet_name}\'!$A{r_no}=TRUE, "{date_sheet_name}", "")'
    ws_log.cell(
        log_idx, 2
    ).value = f'=IF(\'{date_sheet_name}\'!$A{r_no}=TRUE, "{sc_val}", "")'
    ws_log.cell(
        log_idx, 3
    ).value = f'=IF(\'{date_sheet_name}\'!$A{r_no}=TRUE, "{st_name}", "")'
    ws_log.cell(
        log_idx, 4
    ).value = f"=IF('{date_sheet_name}'!$A{r_no}=TRUE, '{date_sheet_name}'!B{r_no}, \"\")"
    ws_log.cell(
        log_idx, 5
    ).value = f"=IF('{date_sheet_name}'!$A{r_no}=TRUE, '{date_sheet_name}'!C{r_no}, \"\")"
    ws_log.cell(
        log_idx, 6
    ).value = f"=IF('{date_sheet_name}'!$A{r_no}=TRUE, {item['old_qty']}, \"\")"

    # G列（確定後）：日付シートの該当セルに直結リンク
    # （※マクロの列並び替えCut&Insertでも自動追従します）
    ws_log.cell(log_idx, 7).value = (
        f"=IF('{date_sheet_name}'!$A{r_no}=TRUE,"
        f" IF('{date_sheet_name}'!{c_let}{r_no}=\"\", 0,"
        f" '{date_sheet_name}'!{c_let}{r_no}), \"\")"
    )

    # H列（増減）：G列 - F列
    if item["old_qty"] == 0:
      ws_log.cell(
          log_idx, 8
      ).value = f"=IF('{date_sheet_name}'!$A{r_no}=TRUE, G{log_idx}, \"\")"
    else:
      ws_log.cell(log_idx, 8).value = (
          f"=IF('{date_sheet_name}'!$A{r_no}=TRUE, G{log_idx}-F{log_idx},"
          ' "")'
      )
    ws_log.cell(log_idx, 8).number_format = DIFF_NUM_FORMAT

  ws_log.freeze_panes = "A2"
  ws_log.auto_filter.ref = f"A1:H{max(ws_log.max_row, 2)}"

  # -------------------------------------------------------------
  # 出荷集約シートの再構築
  # -------------------------------------------------------------
  if "出荷集約" in wb.sheetnames:
    del wb["出荷集約"]
  ws_syukka = wb.create_sheet(title="出荷集約")
  ws_syukka.cell(1, 1).value = "【対象日付】"
  ws_syukka.cell(1, 2).value = f"='{date_sheet_name}'!B1"
  ws_syukka.cell(2, 1).value = "【表示モード】"
  ws_syukka.cell(2, 2).value = "すべて表示"

  dv = DataValidation(
      type="list", formula1='"すべて表示,自社便のみ,佐川便のみ"', allow_blank=True
  )
  ws_syukka.add_data_validation(dv)
  dv.add("B2")

  for c in range(5, total_col_idx):
    ws_syukka.cell(3, c).value = ws.cell(1, c).value
    ws_syukka.cell(4, c).value = ws.cell(2, c).value
  ws_syukka.cell(4, 1).value = "採用"
  ws_syukka.cell(4, 2).value = "コード"
  ws_syukka.cell(4, 3).value = "商品名"
  ws_syukka.cell(4, 4).value = "納品単価"
  ws_syukka.cell(4, total_col_idx).value = "総計"

  for idx, r in enumerate(range(3, current_last_row + 1), start=5):
    ws_syukka.cell(
        idx, 1
    ).value = f"=IF('{date_sheet_name}'!$A{r}=TRUE, TRUE, FALSE)"
    ws_syukka.cell(
        idx, 2
    ).value = f"=IF(A{idx}=TRUE, '{date_sheet_name}'!B{r}, \"\")"
    ws_syukka.cell(
        idx, 3
    ).value = f"=IF(A{idx}=TRUE, '{date_sheet_name}'!C{r}, \"\")"
    ws_syukka.cell(
        idx, 4
    ).value = f"=IF(A{idx}=TRUE, '{date_sheet_name}'!D{r}, \"\")"
    for c in range(5, total_col_idx):
      col_letter = get_column_letter(c)
      sc_key = (
          str(ws.cell(1, c).value).strip() if ws.cell(1, c).value else ""
      )
      dist_target = dist_map.get(sc_key, "自社便のみ")
      ws_syukka.cell(idx, c).value = (
          f'=IF(AND(A{idx}=TRUE, \'{date_sheet_name}\'!$A{r}=TRUE,'
          f' OR($B$2="すべて表示", $B$2="{dist_target}"),'
          f" '{date_sheet_name}'!{col_letter}{r}>0),"
          f" '{date_sheet_name}'!{col_letter}{r}, \"\")"
      )
    ws_syukka.cell(idx, total_col_idx).value = (
        f"=IF(AND(A{idx}=TRUE,"
        f" SUM(E{idx}:{last_store_letter}{idx})>0),"
        f" SUM(E{idx}:{last_store_letter}{idx}), \"\")"
    )

  ws_syukka.freeze_panes = "E5"
  ws_syukka.auto_filter.ref = f"A4:D{current_last_row + 2}"
  ws_syukka.views.sheetView[0].showZeros = False

  # -------------------------------------------------------------
  # 商品別集計シートの再構築
  # -------------------------------------------------------------
  if "商品別集計" in wb.sheetnames:
    del wb["商品別集計"]
  ws_shouhin = wb.create_sheet(title="商品別集計")
  ws_shouhin.cell(1, 1).value = "【対象日付】"
  ws_shouhin.cell(1, 2).value = f"='{date_sheet_name}'!B1"
  ws_shouhin.cell(3, 1).value = "商品コード"
  ws_shouhin.cell(3, 2).value = "商品名"
  ws_shouhin.cell(3, 3).value = "納品単価"
  ws_shouhin.cell(3, 4).value = "確定前総数"
  ws_shouhin.cell(3, 5).value = "確定後総数"
  ws_shouhin.cell(3, 6).value = "差異"

  total_col_letter = get_column_letter(total_col_idx)
  for idx, r in enumerate(range(3, current_last_row + 1), start=4):
    ic_val = ws.cell(r, 2).value
    ic_str = str(int(float(str(ic_val)))) if ic_val else ""
    before_qty = orig_totals.get(ic_str, 0)
    ws_shouhin.cell(
        idx, 1
    ).value = (
        f"=IF('{date_sheet_name}'!$A{r}=TRUE, '{date_sheet_name}'!B{r}, \"\")"
    )
    ws_shouhin.cell(
        idx, 2
    ).value = (
        f"=IF('{date_sheet_name}'!$A{r}=TRUE, '{date_sheet_name}'!C{r}, \"\")"
    )
    ws_shouhin.cell(
        idx, 3
    ).value = (
        f"=IF('{date_sheet_name}'!$A{r}=TRUE, '{date_sheet_name}'!D{r}, \"\")"
    )
    ws_shouhin.cell(
        idx, 4
    ).value = f"=IF('{date_sheet_name}'!$A{r}=TRUE, {before_qty}, \"\")"
    ws_shouhin.cell(idx, 5).value = (
        f"=IF('{date_sheet_name}'!$A{r}=TRUE,"
        f" IF('{date_sheet_name}'!{total_col_letter}{r}=\"\", 0,"
        f" '{date_sheet_name}'!{total_col_letter}{r}), \"\")"
    )
    ws_shouhin.cell(
        idx, 6
    ).value = f"=IF('{date_sheet_name}'!$A{r}=TRUE, E{idx}-D{idx}, \"\")"
    ws_shouhin.cell(idx, 6).number_format = DIFF_NUM_FORMAT

  ws_shouhin.freeze_panes = "A4"
  ws_shouhin.auto_filter.ref = f"A3:F{current_last_row + 1}"
  ws_shouhin.views.sheetView[0].showZeros = False

  # フォント・列幅統一
  for sheet_name in wb.sheetnames:
    target_ws = wb[sheet_name]
    for row in target_ws.iter_rows():
      for cell in row:
        if cell.value is not None:
          is_bold = cell.row in [1, 2, 3, 4] and cell.value in [
              "採用",
              "コード",
              "商品名",
              "納品単価",
              "総計",
              "商品コード",
              "確定前総数",
              "確定後総数",
              "差異",
              "シート",
              "店舗コード",
              "店舗名",
              "修正前(予定)",
              "修正後(確定)",
              "増減",
          ]
          cell.font = MEIRYO_HEADER_FONT if is_bold else MEIRYO_FONT
    auto_fit_columns(target_ws, max_cols=8)

  output = io.BytesIO()
  wb.save(output)
  return output.getvalue(), diff_count, added_items_list, added_stores_list


# UI
col1, col2, col3 = st.columns(3)
with col1:
  f_excel = st.file_uploader("① 発注予定表 (.xls / .xlsx)", type=["xls", "xlsx"])
with col2:
  f_apita = st.file_uploader("② アピタ ピッキング (.pdf)", type=["pdf"])
with col3:
  f_donki = st.file_uploader("③ ドンキ ピッキング (.pdf)", type=["pdf"])

if f_excel and f_apita and f_donki:
  st.markdown("---")
  if st.button(
      "🚀 突合処理を実行してマクロ入り確定版を作成",
      type="primary",
      use_container_width=True,
  ):
    with st.spinner("PDF解析・差分突合・マクロ埋め込み中..."):
      out_bytes, diff_cnt, add_items, add_stores = process_data(
          f_excel, f_apita, f_donki
      )
      st.success("🎉 マクロ入り確定版エクセルの生成が完了しました！")

      m1, m2, m3 = st.columns(3)
      m1.metric("修正セル数（着色・消去含む）", f"{diff_cnt} 件")
      m2.metric("新規追加商品", f"{len(add_items)} 件")
      m3.metric("新規追加店舗", f"{len(add_stores)} 店舗")

      if add_items:
        with st.expander("🆕 追加された商品の詳細"):
          for itm in add_items:
            st.write(f"- {itm}")
      if add_stores:
        with st.expander("🏪 追加された店舗の詳細"):
          for st_name in add_stores:
            st.write(f"- {st_name}")

      st.download_button(
          label="📥 【確定修正版】発注表.xlsm をダウンロード（マクロ有効）",
          data=out_bytes,
          file_name="【確定修正版】発注表.xlsm",
          mime="application/vnd.ms-excel.sheet.macroEnabled.12",
          use_container_width=True,
      )
